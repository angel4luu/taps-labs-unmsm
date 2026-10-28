"""Presentación de evidencia (≤ 15 diapositivas), en el mismo orden que el notebook:
introducción → los 4 componentes (con sus casos límite) → 4 casos de uso completos → conclusiones.

Se arma en .pptx con python-pptx y se exporta a PDF con LibreOffice; ambos quedan en entregables/
(G1_task05_capturas.pptx y G1_task05_capturas.pdf). Las capturas y tablas se leen de evidencia/, así el deck se
regenera consistente tras cada ejecución de `python -m src.flujos --capturas`. Uso: python -m src.presentacion
"""

from __future__ import annotations

import io
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from src.config import DIR_CAPTURAS, DIR_ENTREGABLES, DIR_EVALUACION, DIR_EVIDENCIA, DIR_FLUJOS, RAIZ
from src import flujos

NOMBRE_PPTX = "G1_task05_capturas.pptx"
MAX_DIAPOSITIVAS = 15

# Diseño minimalista: blanco + texto oscuro, un acento guinda (color institucional de San Marcos)
ACENTO = RGBColor.from_string("7A1F2B")
TEXTO = RGBColor.from_string("1C1C1C")
SECUNDARIO = RGBColor.from_string("5E5E5E")
CLARO = RGBColor.from_string("F3F3F3")
LINEA = RGBColor.from_string("DCDCDC")
BLANCO = RGBColor.from_string("FFFFFF")
COLOR_COMPONENTE = {"Reglas": "6E40C9", "NLP/ML": "0969DA", "Retrieval": "1A7F37", "LLM": "BF3989"}
FUENTE = "Calibri"
W, H, M = 13.333, 7.5, 0.6  # lienzo 16:9 (pulgadas) y margen


# ---------------------------------------------------------------- utilidades de dibujo
def _texto(slide, x, y, w, h, parrafos, size=16, color=TEXTO, bold=False, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP, espacio=0):
    """`parrafos`: str (\\n separa párrafos) o lista de párrafos; cada párrafo puede ser una lista
    de (texto, {size, color, bold})."""
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, parrafo in enumerate(parrafos.split("\n") if isinstance(parrafos, str) else parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(espacio)
        for contenido, opciones in (parrafo if isinstance(parrafo, list) else [(parrafo, {})]):
            run = p.add_run()
            run.text = contenido
            run.font.name = FUENTE
            run.font.size = Pt(opciones.get("size", size))
            run.font.bold = opciones.get("bold", bold)
            run.font.color.rgb = opciones.get("color", color)
    return caja


def _sin_estilo(forma):
    estilo = forma._element.find(qn("p:style"))
    if estilo is not None:
        forma._element.remove(estilo)
    return forma


def _forma(slide, tipo, x, y, w, h, relleno, radio=None):
    forma = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    forma.fill.solid()
    forma.fill.fore_color.rgb = relleno
    forma.line.fill.background()
    if radio is not None:
        forma.adjustments[0] = radio
    return _sin_estilo(forma)


def _tarjeta(slide, x, y, w, h, relleno=CLARO):
    return _forma(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, relleno, radio=0.06)


def _imagen(slide, ruta: Path, x, y, max_w, max_h):
    """Inserta la imagen ajustada a la caja (sin deformar) y devuelve (ancho, alto) en pulgadas."""
    with Image.open(ruta) as im:
        im = im.convert("RGB")
        escala = min(max_w / im.width, max_h / im.height)
        w, h = im.width * escala, im.height * escala
        if im.width > 2200:
            im = im.resize((2200, round(im.height * 2200 / im.width)), Image.LANCZOS)
        buffer = io.BytesIO()
        im.save(buffer, "PNG", optimize=True)
    buffer.seek(0)
    foto = slide.shapes.add_picture(buffer, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    foto.line.color.rgb = LINEA
    foto.line.width = Pt(0.75)
    return w, h


def _titulo(slide, texto, subtitulo=None):
    # El tamaño se reduce con el largo para que el título nunca baje a la línea del subtítulo
    size = 28 if len(texto) <= 58 else 24 if len(texto) <= 68 else 21
    _texto(slide, M, 0.4, W - 2 * M, 0.7, texto, size=size, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitulo:
        _texto(slide, M, 1.08, W - 2 * M, 0.5, subtitulo, size=14, color=SECUNDARIO)


def _pie(slide, numero):
    _texto(slide, M, 7.05, 8, 0.3, "Grupo 1 · Tarea 05 · Chatbot híbrido — Asistente académico UNMSM", size=10,
           color=SECUNDARIO)
    _texto(slide, W - M - 1, 7.05, 1, 0.3, str(numero), size=10, color=SECUNDARIO, align=PP_ALIGN.RIGHT)


def _tabla(slide, x, y, anchos, filas, size=11, alto_fila=0.36, colores=None):
    tabla = slide.shapes.add_table(len(filas), len(anchos), Inches(x), Inches(y), Inches(sum(anchos)),
                                   Inches(alto_fila * len(filas))).table
    tabla.first_row = True
    tabla.horz_banding = False
    for j, ancho in enumerate(anchos):
        tabla.columns[j].width = Inches(ancho)
    for i, fila in enumerate(filas):
        tabla.rows[i].height = Inches(alto_fila)
        for j, valor in enumerate(fila):
            celda = tabla.cell(i, j)
            celda.fill.solid()
            celda.fill.fore_color.rgb = CLARO if i == 0 else BLANCO
            celda.margin_left = celda.margin_right = Inches(0.06)
            celda.margin_top = celda.margin_bottom = Inches(0.03)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = celda.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = str(valor)
            run.font.name, run.font.size, run.font.bold = FUENTE, Pt(size), i == 0
            color = (colores or {}).get((i, j))
            run.font.color.rgb = RGBColor.from_string(color) if color else TEXTO
    return tabla


def _nueva(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------- diapositivas
def _portada(prs):
    s = _nueva(prs)
    _forma(s, MSO_SHAPE.RECTANGLE, 0, 0, 0.35, H, ACENTO)
    _texto(s, 1.0, 1.4, 11, 0.5, "TAREA 05 · CHATBOT HÍBRIDO", size=16, bold=True, color=ACENTO)
    _texto(s, 1.0, 1.95, 11.3, 1.6, "Asistente académico de pregrado UNMSM", size=44, bold=True)
    _texto(s, 1.0, 3.35, 11, 0.9, "Reglas · Recuperación de información · NLP/ML · LLM — sobre el Reglamento "
                                  "General de Matrícula de Pregrado 2026 y la Ley Universitaria 30220",
           size=18, color=SECUNDARIO)
    integrantes = ["Ayala Alberca, Sebastian", "Gamarra Chavez, Diego Valentino", "Reyes Fabian, Jhan Brayan Zamir",
                   "Jara Espinoza, Angela Lucia"]
    _texto(s, 1.0, 4.75, 6, 1.6, [[("Grupo 1", {"bold": True, "size": 15})]] + integrantes, size=14, espacio=2)
    _texto(s, 7.4, 4.75, 5.2, 1.6, "Taller de Aplicaciones Sociales · UNMSM\nEntregables: notebook "
                                   "G1_task05_chatbot_hibrido.ipynb y esta presentación", size=13, color=SECUNDARIO)


def _arquitectura(prs, n):
    s = _nueva(prs)
    _titulo(s, "Introducción: un chatbot híbrido para trámites de pregrado",
            "Cada mensaje recorre los cuatro componentes en orden; el primero que puede responder, responde. Cada "
            "respuesta muestra qué componente la produjo.")
    etapas = [
        ("Reglas", "Saludo, ayuda, nombre, despedida y guardarraíles (deshonestidad, datos personales, "
                   "inyección)."),
        ("NLP/ML", "Intención (TF-IDF + NB/LR/SVM vs. embeddings; se usa A+B) + entidades (regex, listas, "
                   "spaCy) + política: 3 bandas, pedir detalle, memoria."),
        ("Retrieval", "TF-IDF + embeddings sobre 41 FAQ y 60 pasajes oficiales, filtrado por la intención; "
                      "devuelve el texto tal cual, con fuente."),
        ("LLM", "Gemini con los pasajes recuperados (RAG): respuesta en JSON con citas validadas; sin conexión, "
                "cita el artículo textual."),
    ]
    ancho, y = 2.85, 1.85
    for k, (nombre, texto) in enumerate(etapas):
        x = M + k * (ancho + 0.25)
        color = RGBColor.from_string(COLOR_COMPONENTE[nombre])
        _forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, ancho, 0.55, color, radio=0.15)
        _texto(s, x, y, ancho, 0.55, f"{k + 1}. {nombre}", size=17, bold=True, color=BLANCO,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _tarjeta(s, x, y + 0.7, ancho, 2.0)
        _texto(s, x + 0.15, y + 0.85, ancho - 0.3, 1.8, texto, size=14)
        if k < 3:
            _texto(s, x + ancho, y + 0.05, 0.25, 0.45, "›", size=24, bold=True, color=SECUNDARIO,
                   align=PP_ALIGN.CENTER)
    widget = DIR_CAPTURAS / "widget_jupyterlab.png"
    if widget.exists():
        _texto(s, M, 4.75, 4.4, 0.4, "El chat dentro del notebook (ipywidgets)", size=13, bold=True)
        _imagen(s, widget, M, 5.15, 6.0, 1.75)
        x_texto = 7.0
    else:
        x_texto = M
    _texto(s, x_texto, 4.85, W - x_texto - M, 2.1, [
        [("Políticas visibles en cada respuesta: ", {"bold": True})],
        "Pide detalle · Confirma tema · Memoria de sesión · Rechazo · Fuera de dominio · Sin filtro · "
        "Fallback offline",
        [("Fuentes: ", {"bold": True}),
         ("Reglamento General de Matrícula Pregrado 2026 (RR N°002642-2026) y Ley Universitaria 30220 Art. 102; "
          "FAQ del equipo marcadas como referenciales.", {})],
    ], size=14, espacio=6)
    _pie(s, n)


def _diagrama(prs, n):
    """Flujo de un mensaje (diagrama Mermaid de src/diagrama.py) con una leyenda por componente."""
    s = _nueva(prs)
    _titulo(s, "Flujo de un mensaje",
            "Las preguntas de cada paso deciden qué componente responde; el primero que puede, responde.")
    _imagen(s, DIR_EVIDENCIA / "flujo_mensaje.png", M, 1.6, W - 2 * M, 3.4)
    leyenda = [
        ("Reglas", "Patrones fijos: saludo, ayuda, nombre, despedida. Los pedidos indebidos se rechazan aquí."),
        ("NLP/ML", "Detecta la intención y los datos. Si duda, pregunta «¿X o Y?»; si falta un dato (la "
                   "facultad), lo pide y lo recuerda."),
        ("Retrieval", "Busca en las FAQ y el reglamento dentro del tema detectado; una pregunta general recibe la "
                      "FAQ tal cual, con su fuente."),
        ("LLM", "Preguntas personales o de seguimiento: Gemini redacta con los artículos recuperados y los cita."),
    ]
    ancho, y = (W - 2 * M - 3 * 0.25) / 4, 5.2
    for k, (nombre, texto) in enumerate(leyenda):
        x = M + k * (ancho + 0.25)
        _forma(s, MSO_SHAPE.RECTANGLE, x, y, 0.07, 1.55, RGBColor.from_string(COLOR_COMPONENTE[nombre]))
        _texto(s, x + 0.18, y, ancho - 0.2, 1.6, [[(nombre, {"bold": True, "size": 14})], texto], size=12, espacio=3)
    _pie(s, n)


def _galeria(slide, capturas, x, y, ancho, alto, separacion=0.15):
    """Capturas lado a lado (las conversaciones largas vienen partidas en grupos de turnos)."""
    if not capturas:
        return 0
    ancho_cada = (ancho - separacion * (len(capturas) - 1)) / len(capturas)
    usado = 0
    for k, captura in enumerate(capturas):
        w, _ = _imagen(slide, captura, x + usado, y, ancho_cada, alto)
        usado += w + separacion
    return usado - separacion


def _componente(prs, n, nombre, titulo, subtitulo, puntos, capturas, nota=None):
    """Diapositiva de un componente: qué hace (izquierda) y cómo se ve en el chat (derecha)."""
    s = _nueva(prs)
    _titulo(s, titulo, subtitulo)
    color = RGBColor.from_string(COLOR_COMPONENTE[nombre])
    _forma(s, MSO_SHAPE.RECTANGLE, M, 1.72, 0.08, 3.4, color)
    _texto(s, M + 0.3, 1.7, 4.9, 3.6, [[("Cómo funciona", {"bold": True, "size": 15})]] + [f"• {p}" for p in puntos],
           size=13.5, espacio=7)
    if nota:
        _tarjeta(s, M, 5.45, 5.2, 1.35)
        _texto(s, M + 0.18, 5.55, 4.85, 1.2, nota, size=12.5)
    _galeria(s, capturas, 6.15, 1.65, W - 6.15 - M, 5.25)
    _pie(s, n)


def _traza(slide, traza, x, y, ancho):
    """Franja con un chip por turno: «T1 Reglas», «T2 NLP/ML · Pide detalle»…"""
    _texto(slide, x, y, 1.0, 0.36, "Turnos:", size=12, bold=True, color=SECUNDARIO, anchor=MSO_ANCHOR.MIDDLE)
    cx = x + 0.85
    for fila in traza:
        etiqueta = f"T{fila['Turno']} {fila['Componente']}" + (f" · {fila['Política']}" if fila["Política"] != "—" else "")
        w = 0.22 + 0.078 * len(etiqueta)
        if cx + w > x + ancho:
            break
        base = fila["Componente"].split(" ")[0]
        color = COLOR_COMPONENTE.get(base, "9A6700" if "local" in fila["Componente"] else "57606A")
        _forma(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx, y, w, 0.36, RGBColor.from_string(color), radio=0.3)
        _texto(slide, cx, y, w, 0.36, etiqueta, size=11, bold=True, color=BLANCO, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE)
        cx += w + 0.12


def _flujo(prs, n, datos, traza, capturas, parte=None):
    s = _nueva(prs)
    sufijo = f" (parte {parte})" if parte else ""
    _titulo(s, f"{flujos.titulo(datos)}{sufijo}", datos["historia"])
    _traza(s, traza, M, 1.6, W - 2 * M)
    _galeria(s, capturas, M, 2.15, W - 2 * M, 4.75, separacion=0.2)
    _pie(s, n)


def _evaluacion_cv(prs, n, rep):
    s = _nueva(prs)
    sel = rep["seleccion"]
    _titulo(s, f"Componente 3 · NLP/ML (1/3): gana {sel['ganador']}",
            "300 frases etiquetadas (10 intenciones). Validación cruzada de 5 folds sobre las 250 de entrenamiento.")
    filas = [["Candidato", "Rama", "F1 macro (media ± std)"]]
    ramas = {"A·NB": "Léxica: TF-IDF + Naive Bayes", "A·LR": "Léxica: TF-IDF + Reg. Logística",
             "A·SVM": "Léxica: TF-IDF + SVM calibrada", "B·Emb+LR": "Semántica: embeddings + Reg. Logística",
             "A+B promedio": "Combinación: promedio de probabilidades"}
    for c in rep["cv"]:
        filas.append([c["candidato"], ramas.get(c["candidato"], ""), f"{c['f1_macro_media']:.3f} ± "
                                                                      f"{c['f1_macro_std']:.3f}"])
    _tabla(s, M, 1.75, [1.45, 3.4, 1.75], filas, size=12, alto_fila=0.42)
    _tarjeta(s, M, 4.5, 6.6, 1.55)
    _texto(s, M + 0.2, 4.62, 6.2, 1.35, [[("Regla de selección (fijada antes de ver resultados)", {"bold": True})],
                                         sel["regla_aplicada"]], size=11.5, espacio=4)
    _imagen(s, DIR_EVALUACION / "comparacion_cv.png", 7.55, 1.75, W - 7.55 - M, 5.0)
    _pie(s, n)


def _evaluacion_prueba(prs, n, rep):
    s = _nueva(prs)
    t = rep["test"]
    _titulo(s, f"Componente 3 · NLP/ML (2/3): resultados en prueba",
            f"{t['n']} frases separadas desde el inicio: nunca se usaron para entrenar ni para elegir el modelo.")
    _imagen(s, DIR_EVALUACION / "matriz_confusion.png", M, 1.85, 5.6, 5.0)
    filas = [["Métrica", "Resultado", "Objetivo"],
             ["Accuracy", f"{t['accuracy']:.3f}", "≥ 0.85"],
             ["F1 macro", f"{t['f1_macro']:.3f}", "≥ 0.80"]]
    _tabla(s, 6.7, 1.85, [3.2, 1.25, 1.2], filas, size=12, alto_fila=0.42)
    errores = rep.get("errores", [])[:4]
    _texto(s, 6.7, 3.4, W - 6.7 - M, 3.4,
           [[("Errores típicos (temas vecinos, con confianza baja → el chatbot pide confirmar):", {"bold": True})]] +
           [f"«{e['texto']}» → {e['predicho']} (era {e['real']}, {e['confianza']:.2f})" for e in errores],
           size=11.5, espacio=3)
    _pie(s, n)


def _conclusiones(prs, n, rep):
    s = _nueva(prs)
    _titulo(s, "Resultados y conclusiones")
    t = rep["test"]
    pipeline = rep.get("retrieval", {}).get("prueba_pipeline", {})
    puntos = [
        ("Reglas", "Guion de saludo, ayuda, nombre, despedida y guardarraíles: 100 % de respuestas esperadas."),
        ("NLP/ML", f"La combinación léxica + semántica supera a cada rama sola (F1 CV {rep['cv'][-1]['f1_macro_media']:.2f}); "
                   f"en prueba accuracy {t['accuracy']:.2f} y F1 {t['f1_macro']:.2f}. La política de bandas convierte "
                   f"las dudas en preguntas de confirmación en vez de respuestas equivocadas."),
        ("Retrieval", f"Con el filtro de intención, acierto@1 {pipeline.get('hit_at_1', 0):.2f} y no-respuesta "
                      f"{pipeline.get('no_respuesta_correcta', 0):.2f} en preguntas nuevas; respuestas textuales con fuente."),
        ("LLM", "Solo redacta con los pasajes recuperados y cita artículos; sin clave o sin conexión el chat sigue "
                "funcionando: cita el artículo textual y remite a la oficina."),
        ("Límites", "Datos de intención escritos por el equipo (300 frases); fechas y montos no se inventan: se remite "
                    "a la oficina oficial. El Reglamento de Grados y Títulos no está publicado en texto completo."),
    ]
    y = 1.4
    for nombre, texto in puntos:
        color = RGBColor.from_string(COLOR_COMPONENTE.get(nombre, "57606A"))
        _forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, M, y, 1.7, 0.85, color, radio=0.15)
        _texto(s, M, y, 1.7, 0.85, nombre, size=15, bold=True, color=BLANCO, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE)
        _texto(s, M + 1.95, y + 0.05, W - 2 * M - 1.95, 0.8, texto, size=13.5, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.08
    _pie(s, n)


def convertir_a_pdf(pptx: Path, carpeta: Path) -> Path:
    """Exporta el .pptx a PDF con LibreOffice headless (perfil temporal: funciona aunque esté abierto)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("Se necesita LibreOffice (comando 'soffice') para exportar a PDF")
    carpeta.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as perfil:
        subprocess.run([soffice, f"-env:UserInstallation={Path(perfil).as_uri()}", "--headless", "--convert-to",
                        "pdf", "--outdir", str(carpeta), str(pptx)], check=True, capture_output=True, timeout=300)
    pdf = carpeta / f"{pptx.stem}.pdf"
    if not pdf.exists():
        raise RuntimeError(f"LibreOffice no generó {pdf.name}")
    return pdf


def generar() -> Path:
    """Arma la presentación (.pptx) y la exporta a PDF; ambos archivos quedan en entregables/."""
    rep = json.loads((DIR_EVALUACION / "reporte.json").read_text(encoding="utf-8"))
    trazas = json.loads((DIR_FLUJOS / "trazas.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Tarea 05 · Chatbot híbrido — Asistente académico UNMSM"
    prs.core_properties.author = "Grupo 1"

    _portada(prs)
    _arquitectura(prs, 2)
    _diagrama(prs, 3)
    cap = lambda nombre: sorted(DIR_CAPTURAS.glob(f"escena_{nombre}_[0-9]*.png"))
    r = rep.get("retrieval", {}).get("prueba_pipeline", {})
    _componente(prs, 4, "Reglas", "Componente 1 · Basado en reglas",
                "Respuestas fijas e inmediatas para lo predecible, y protección antes que todo lo demás.",
                ["Expresiones regulares sobre el texto normalizado (data/reglas/patrones.yaml).",
                 "Conversación: saludo, menú de ayuda, nombre del estudiante (se recuerda), gracias, despedida.",
                 "Protección: pedidos deshonestos, datos personales de terceros e intentos de cambiar las "
                 "instrucciones se rechazan antes de llegar a otro componente.",
                 "Si ninguna regla coincide, el mensaje pasa al siguiente componente."],
                cap("reglas"))
    _componente(prs, 5, "Retrieval", "Componente 2 · Basado en recuperación de información",
                "Respuestas verificables: se devuelve el texto existente, con su fuente.",
                ["Base: 41 preguntas frecuentes y 60 pasajes del Reglamento de Matrícula 2026 (uno por artículo).",
                 "Similitud de coseno combinando TF-IDF (palabras) y embeddings (significado: reconoce sinónimos).",
                 "La búsqueda se limita al tema que detectó el componente NLP/ML.",
                 "Si la FAQ más parecida supera el umbral, se responde tal cual; si no, los artículos pasan al LLM."],
                cap("recuperacion"),
                nota=f"En 53 preguntas nuevas: respuesta correcta en el {r.get('hit_at_1', 0):.0%}; preguntas "
                     f"ajenas correctamente no respondidas: {r.get('no_respuesta_correcta', 0):.0%}.")
    _evaluacion_cv(prs, 6, rep)
    _evaluacion_prueba(prs, 7, rep)
    u = rep["umbrales"]
    _componente(prs, 8, "NLP/ML", "Componente 3 · NLP/ML (3/3): entidades y decisión",
                "Con la intención y los datos del mensaje, el chatbot decide qué hacer.",
                ["Entidades: expresiones regulares (periodo, «dos veces»), listas del dominio (20 facultades con "
                 "siglas, cursos, trámites) y el reconocedor de spaCy.",
                 f"Confianza ≥ {u['aceptar']}: acepta el tema y filtra la búsqueda.",
                 f"Entre {u['rechazar']} y {u['aceptar']}: pregunta «¿es sobre X o sobre Y?».",
                 f"Menor a {u['rechazar']}: busca sin filtro de tema. Lo ajeno a la UNMSM se rechaza.",
                 "Si falta un dato necesario (p. ej. la facultad), lo pide una vez y lo recuerda."],
                cap("nlp"))
    _componente(prs, 9, "LLM", "Componente 4 · Basado en LLM",
                "Redacta respuestas a preguntas abiertas usando solo el reglamento, con citas.",
                ["Gemini recibe los artículos recuperados, el tema, los datos del estudiante y la conversación previa.",
                 "Instrucciones fijas: solo usar esos artículos, citarlos, no inventar fechas ni montos, decir «no lo "
                 "sé» si no alcanzan.",
                 "Las citas se verifican contra los artículos enviados.",
                 "Sin clave o sin conexión: cita el artículo textual y remite a la oficina."],
                cap("llm") + cap("llm_offline"))
    n = 10
    for datos in flujos.cargar(tipo="caso"):
        capturas = sorted(DIR_CAPTURAS.glob(f"{flujos.nombre_archivo(datos)}_*.png"))
        if not capturas:
            raise FileNotFoundError(f"Faltan capturas del caso {datos['id']} (python -m src.flujos --capturas)")
        traza = trazas[str(datos["id"])]
        grupos = [capturas[k:k + 3] for k in range(0, len(capturas), 3)]  # hasta 3 columnas por diapositiva
        for k, grupo in enumerate(grupos, start=1):
            _flujo(prs, n, datos, traza, grupo, parte=k if len(grupos) > 1 else None)
            n += 1
    _conclusiones(prs, n, rep)
    total = len(prs.slides)
    if total > MAX_DIAPOSITIVAS:
        raise RuntimeError(f"La presentación tiene {total} diapositivas (máximo {MAX_DIAPOSITIVAS})")

    DIR_ENTREGABLES.mkdir(parents=True, exist_ok=True)
    pptx = DIR_ENTREGABLES / NOMBRE_PPTX  # entregable editable, junto a su versión PDF
    prs.save(pptx)
    return convertir_a_pdf(pptx, DIR_ENTREGABLES)


if __name__ == "__main__":
    pdf = generar()
    for archivo in (pdf.with_suffix(".pptx"), pdf):
        print(f"✅ {archivo.relative_to(RAIZ)} ({archivo.stat().st_size / 1024 ** 2:.1f} MB)")
