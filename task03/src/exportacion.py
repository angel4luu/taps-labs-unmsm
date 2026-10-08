"""Entregables empaquetados: datos (CSV + XLSX + LEEME.txt) y código del proyecto, cada uno en su .zip."""
import io
import logging
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import config as cfg
from src.analisis import resumen

log = logging.getLogger(__name__)

NOMBRE_CSV = f"{cfg.PREFIJO}_resenas_rappi.csv"
NOMBRE_XLSX = f"{cfg.PREFIJO}_resenas_rappi.xlsx"
NOMBRE_LEEME = "LEEME.txt"
NOMBRE_ZIP = f"{cfg.PREFIJO}_data.zip"
HOJAS = ["Reseñas", "Diccionario", "Resumen"]

# (columna, tipo, descripción)
DICCIONARIO = [
    ("id_resena", "texto (40)", "Identificador único: hash SHA-1 de usuario + fecha + comentario"),
    ("usuario", "texto", "Nombre público del autor en Google Play"),
    ("calificacion", "entero 1–5", "Estrellas otorgadas por el usuario"),
    ("fecha", "fecha", "Fecha de publicación de la reseña (o de su última edición)"),
    ("comentario", "texto", "Texto de la reseña (máximo 500 caracteres en Google Play)"),
    ("longitud_comentario", "entero", "Cantidad de caracteres del comentario"),
    ("votos_utiles", "entero", "Personas que marcaron la reseña como útil (0 si ninguna)"),
    ("tiene_respuesta", "booleano", "True si Rappi respondió la reseña"),
    ("respuesta_empresa", "texto / vacío", "Texto de la respuesta de Rappi; vacío si no respondió"),
    ("fecha_respuesta", "fecha / vacío", "Fecha de la respuesta de Rappi"),
    ("dias_respuesta", "entero / vacío",
     "Días entre la reseña y la respuesta; vacío si no hay respuesta o si la reseña fue editada después"),
    ("editada_tras_respuesta", "booleano",
     "True si el usuario editó su reseña después de la respuesta (ver nota en LEEME.txt)"),
    ("app", "texto", "Aplicación de origen (Rappi)"),
    ("fecha_extraccion", "fecha y hora", "Momento en que el scraper extrajo la reseña"),
]

NOTA_EDITADAS = (
    "En Google Play, cuando un usuario edita su reseña, esta toma la fecha de la edición, pero la "
    "respuesta de Rappi conserva su fecha original. Por eso, en las reseñas con "
    "editada_tras_respuesta = True la respuesta parece anterior a la reseña y su dias_respuesta "
    "queda vacío (no es medible)."
)


def diccionario(df: pd.DataFrame) -> pd.DataFrame:
    """Diccionario de datos con un ejemplo real de cada columna."""
    filas = []
    for columna, tipo, descripcion in DICCIONARIO:
        valores = df[columna].dropna()
        ejemplo = valores.iloc[0] if len(valores) else ""
        if isinstance(ejemplo, pd.Timestamp):
            ejemplo = ejemplo.strftime("%Y-%m-%d %H:%M" if columna == "fecha_extraccion" else "%Y-%m-%d")
        ejemplo = str(ejemplo)[:60]
        filas.append({"columna": columna, "tipo": tipo, "descripcion": descripcion, "ejemplo": ejemplo})
    return pd.DataFrame(filas)


def metadatos(df: pd.DataFrame) -> dict:
    """Fuente, método y cifras clave (hoja Resumen y LEEME.txt)."""
    return {
        "Grupo": cfg.GRUPO,
        "Aplicación": f"{cfg.APP_NOMBRE} ({cfg.APP_ID})",
        "Fuente": cfg.URL,
        "Método": "Selenium + ChromeDriver (Google Chrome automatizado)",
        "Orden de extracción": cfg.ORDEN,
        "Fecha de extracción": f"{df['fecha_extraccion'].min():%d/%m/%Y %H:%M}",
        **resumen(df),
    }


def _hoja_resenas(ws, df: pd.DataFrame):
    """Encabezado en negrita y fijo, filtros, anchos, texto ajustado y fechas reales."""
    anchos = {"id_resena": 14, "usuario": 24, "comentario": 70, "respuesta_empresa": 70,
              "fecha": 12, "fecha_respuesta": 14, "fecha_extraccion": 18}
    formatos = {"fecha": "dd/mm/yyyy", "fecha_respuesta": "dd/mm/yyyy", "fecha_extraccion": "dd/mm/yyyy hh:mm"}
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="E8E8E4")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, columna in enumerate(df.columns, start=1):
        letra = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letra].width = anchos.get(columna, max(len(columna) + 2, 10))
        ajustar = columna in ("comentario", "respuesta_empresa")
        for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            celda.alignment = Alignment(vertical="top", wrap_text=ajustar)
            if columna in formatos:
                celda.number_format = formatos[columna]


def _hoja_simple(ws, anchos: list[int]):
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="E8E8E4")
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)


def exportar_xlsx(df: pd.DataFrame, ruta: Path) -> Path:
    meta = pd.DataFrame(list(metadatos(df).items()), columns=["campo", "valor"])
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=HOJAS[0], index=False)
        diccionario(df).to_excel(writer, sheet_name=HOJAS[1], index=False)
        meta.to_excel(writer, sheet_name=HOJAS[2], index=False)
        _hoja_resenas(writer.sheets[HOJAS[0]], df)
        _hoja_simple(writer.sheets[HOJAS[1]], [24, 16, 70, 40])
        _hoja_simple(writer.sheets[HOJAS[2]], [30, 80])
    return ruta


def exportar_leeme(df: pd.DataFrame, ruta: Path) -> Path:
    lineas = [f"{cfg.PREFIJO.upper()} — Reseñas de {cfg.APP_NOMBRE} en Google Play", "=" * 60, ""]
    lineas += [f"{campo}: {valor}" for campo, valor in metadatos(df).items()]
    lineas += ["", "ARCHIVOS", "-" * 60,
               f"{NOMBRE_CSV}   datos en CSV (UTF-8, separador coma, fechas AAAA-MM-DD)",
               f"{NOMBRE_XLSX}  mismos datos en Excel + hojas Diccionario y Resumen",
               "", "DICCIONARIO DE DATOS", "-" * 60]
    for columna, tipo, descripcion in DICCIONARIO:
        lineas.append(f"{columna:<24}{tipo:<16}{descripcion}")
    lineas += ["", "NOTA SOBRE RESEÑAS EDITADAS", "-" * 60, NOTA_EDITADAS, "",
               "Uso académico. Datos públicos extraídos sin inicio de sesión, con pausas entre solicitudes."]
    ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8-sig")
    return ruta


def verificar_zip(ruta_zip: Path, df: pd.DataFrame) -> dict:
    """Reabre el zip y comprueba integridad, contenido, filas e ids. Lanza AssertionError si falla."""
    with zipfile.ZipFile(ruta_zip) as zf:
        assert zf.testzip() is None, "zip corrupto"
        nombres = sorted(zf.namelist())
        assert nombres == sorted([NOMBRE_CSV, NOMBRE_XLSX, NOMBRE_LEEME]), f"contenido inesperado: {nombres}"
        csv = pd.read_csv(zf.open(NOMBRE_CSV), encoding="utf-8-sig")
        assert len(csv) == len(df), f"el CSV tiene {len(csv)} filas, se esperaban {len(df)}"
        assert set(csv["id_resena"]) == set(df["id_resena"]), "los ids del CSV no coinciden"
        libro = load_workbook(io.BytesIO(zf.read(NOMBRE_XLSX)), read_only=True)
        assert libro.sheetnames == HOJAS, f"hojas inesperadas: {libro.sheetnames}"
        filas_xlsx = libro[HOJAS[0]].max_row - 1
        assert filas_xlsx == len(df), f"el XLSX tiene {filas_xlsx} filas, se esperaban {len(df)}"
        archivos = {i.filename: f"{i.file_size / 1024:,.0f} KB" for i in zf.infolist()}
    log.info("Zip verificado: %s", archivos)
    return {"zip": ruta_zip.name, "tamaño": f"{ruta_zip.stat().st_size / 1024:,.0f} KB",
            "filas": len(df), "archivos": archivos}


def exportar_y_empaquetar(df: pd.DataFrame) -> dict:
    """Exporta CSV, XLSX y LEEME.txt, crea el zip en entregables/ y lo verifica."""
    cfg.DIR_PROCESADA.mkdir(parents=True, exist_ok=True)
    cfg.DIR_ENTREGABLES.mkdir(parents=True, exist_ok=True)
    archivos = [cfg.DIR_PROCESADA / nombre for nombre in (NOMBRE_CSV, NOMBRE_XLSX, NOMBRE_LEEME)]
    csv = df.copy()
    for columna in ("fecha", "fecha_respuesta"):  # solo fecha; fecha_extraccion conserva la hora
        csv[columna] = csv[columna].dt.strftime("%Y-%m-%d")
    csv.to_csv(archivos[0], index=False, encoding="utf-8-sig")
    exportar_xlsx(df, archivos[1])
    exportar_leeme(df, archivos[2])

    ruta_zip = cfg.DIR_ENTREGABLES / NOMBRE_ZIP
    with zipfile.ZipFile(ruta_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for archivo in archivos:
            zf.write(archivo, arcname=archivo.name)  # en la raíz del zip, sin carpetas
    return verificar_zip(ruta_zip, df)


# ---------------------------------------------------------------- código del proyecto
NOMBRE_ZIP_CODIGO = f"{cfg.PREFIJO}_codigo.zip"
ARCHIVOS_CODIGO = [f"{cfg.PREFIJO}_scraping_rappi.ipynb", "prueba_rapida.py", "inspeccion.py",
                   "environment.yml", "README.md", "scraping.log"]
CARPETAS_CODIGO = ["src", "tests"]


def empaquetar_codigo() -> dict:
    """Crea G1_task03_codigo.zip con el proyecto ejecutable (sin datos, capturas ni temporales)
    y lo verifica: lo descomprime en una carpeta vacía y corre allí las pruebas."""
    archivos = [cfg.BASE / nombre for nombre in ARCHIVOS_CODIGO]
    for carpeta in CARPETAS_CODIGO:
        archivos += sorted(p for p in (cfg.BASE / carpeta).rglob("*.py") if "__pycache__" not in p.parts)
    faltantes = [p.name for p in archivos if not p.exists()]
    assert not faltantes, f"faltan archivos: {faltantes}"

    ruta_zip = cfg.DIR_ENTREGABLES / NOMBRE_ZIP_CODIGO
    with zipfile.ZipFile(ruta_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for archivo in archivos:
            zf.write(archivo, arcname=archivo.relative_to(cfg.BASE))

    with tempfile.TemporaryDirectory() as carpeta, zipfile.ZipFile(ruta_zip) as zf:
        assert zf.testzip() is None, "zip corrupto"
        zf.extractall(carpeta)
        prueba = subprocess.run([sys.executable, "-m", "tests.test_parser"], cwd=carpeta,
                                capture_output=True, text=True)
        assert prueba.returncode == 0, f"las pruebas fallan en el zip:\n{prueba.stderr}"
        importa = subprocess.run([sys.executable, "-c", "import src.scraper, src.presentacion"], cwd=carpeta,
                                 capture_output=True, text=True)
        assert importa.returncode == 0, f"el código del zip no importa:\n{importa.stderr}"
    log.info("Zip de código verificado: %d archivos", len(archivos))
    return {"zip": ruta_zip.name, "tamaño": f"{ruta_zip.stat().st_size / 1024:,.0f} KB",
            "archivos": len(archivos), "pruebas": prueba.stdout.strip().splitlines()[-1]}
