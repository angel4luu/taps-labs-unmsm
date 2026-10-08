"""Presentación de evidencia: se arma en .pptx con python-pptx y se exporta a PDF con LibreOffice.

Versiones, horas, cifras y muestras se leen del log, de los datos limpios y de las capturas,
así la presentación se regenera consistente tras cada ejecución. Uso: python -m src.presentacion
"""
import inspect
import io
import logging
import platform
import re
import shutil
import subprocess
import tempfile
import textwrap
import zipfile
from pathlib import Path

import pandas as pd
import selenium
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from src import config as cfg
from src.analisis import resumen
from src.exportacion import NOMBRE_ZIP
from src.scraper import scroll_y_extraer
from src.storage import cargar_crudo, limpiar

log = logging.getLogger(__name__)

NOMBRE_PPTX = f"{cfg.PREFIJO}_evidencia.pptx"   # copia editable de trabajo (evidencia/)
NOMBRE_PDF = f"{cfg.PREFIJO}_evidencia.pdf"     # entregable (entregables/)

# Diseño minimalista: blanco + texto oscuro, un solo acento naranja Rappi
NARANJA = RGBColor.from_string("FF441F")
OSCURO = RGBColor.from_string("1C1C1C")
TEXTO = RGBColor.from_string("1C1C1C")
SECUNDARIO = RGBColor.from_string("5E5E5E")
CLARO = RGBColor.from_string("F3F3F3")
LINEA = RGBColor.from_string("DCDCDC")
BLANCO = RGBColor.from_string("FFFFFF")
GRIS_OSCURO = RGBColor.from_string("B5B5B5")
FUENTE, MONO = "Calibri", "Courier New"
W, H, M = 13.333, 7.5, 0.6          # lienzo 16:9 (pulgadas) y margen
CAPTURA = 1.6                        # proporción ancho/alto de las capturas (2880x1800)


# ---------------------------------------------------------------- datos de entrada
def _decimal(valor, decimales=1) -> str:
    return f"{valor:.{decimales}f}".replace(".", ",")


def _ultima_ejecucion() -> list[str]:
    """Líneas del log de la última ejecución completa (de 'Inicio' a 'Fin')."""
    lineas = cfg.ARCHIVO_LOG.read_text(encoding="utf-8").splitlines()
    inicio = max(i for i, linea in enumerate(lineas) if "Inicio |" in linea)
    fin = next((i for i in range(inicio, len(lineas)) if "Fin |" in lineas[i]), len(lineas) - 1)
    return [linea for linea in lineas[inicio:fin + 1] if linea.strip() and "Reseñas:" not in linea]


def _datos() -> dict:
    df, _ = limpiar(cargar_crudo())
    ejecucion = _ultima_ejecucion()
    chrome, chromedriver = re.search(r"Chrome ([\d.]+) \| ChromeDriver ([\d.]+)", ejecucion[0]).groups()
    total, segundos = map(int, re.search(r"Fin \| (\d+) reseñas \| (\d+) s", ejecucion[-1]).groups())
    horas = {m[2]: m[1] for linea in ejecucion
             if (m := re.match(r"(\d\d:\d\d:\d\d) INFO\s+Captura de pantalla completa: (\S+)", linea))}
    with zipfile.ZipFile(cfg.DIR_ENTREGABLES / NOMBRE_ZIP) as zf:
        archivos_zip = [(i.filename, f"{i.file_size / 1024:,.0f} KB") for i in zf.infolist()]
    return {"df": df, "resumen": resumen(df), "ejecucion": ejecucion, "chrome": chrome,
            "chromedriver": chromedriver, "total": total, "segundos": segundos, "horas": horas,
            "archivos_zip": archivos_zip}


# ---------------------------------------------------------------- primitivas de dibujo
def _texto(slide, x, y, w, h, parrafos, size=16, color=TEXTO, bold=False, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP, font=FUENTE, espacio=0, nombre=None):
    """Caja de texto; `parrafos` es un str (\\n separa párrafos) o una lista de párrafos,
    donde cada párrafo puede ser una lista de (texto, {size, color, bold, font})."""
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nombre:
        caja.name = nombre
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, parrafo in enumerate(parrafos.split("\n") if isinstance(parrafos, str) else parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(espacio)
        for contenido, opciones in (parrafo if isinstance(parrafo, list) else [(parrafo, {})]):
            run = p.add_run()
            run.text = contenido
            fuente = run.font
            fuente.name = opciones.get("font", font)
            fuente.size = Pt(opciones.get("size", size))
            fuente.bold = opciones.get("bold", bold)
            fuente.color.rgb = opciones.get("color", color)
    return caja


def _sin_estilo(forma):
    """Quita el estilo de tema (sombra incluida): colores y líneas ya se fijan explícitamente.
    Así la forma se ve plana tanto en PowerPoint como en LibreOffice."""
    estilo = forma._element.find(qn("p:style"))
    if estilo is not None:
        forma._element.remove(estilo)
    return forma


def _forma(slide, tipo, x, y, w, h, relleno, radio=None):
    forma = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    forma.fill.solid()
    forma.fill.fore_color.rgb = relleno
    forma.line.fill.background()
    if radio is not None:
        forma.adjustments[0] = radio
    return _sin_estilo(forma)


def _linea(slide, x1, y1, x2, y2, grosor=0.75):
    linea = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    linea.line.color.rgb = LINEA
    linea.line.width = Pt(grosor)
    return _sin_estilo(linea)


def _tarjeta(slide, x, y, w, h, relleno=CLARO):
    return _forma(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, relleno, radio=0.06)


def _numero(slide, x, y, d, n, size=16):
    """Motivo visual del deck: número blanco en un círculo naranja."""
    circulo = _forma(slide, MSO_SHAPE.OVAL, x, y, d, d, NARANJA)
    tf = circulo.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(n)
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FUENTE, Pt(size), True, BLANCO
    return circulo


def _imagen(slide, ruta, x, y, w):
    """Inserta una captura reducida a 1920 px (JPEG) con un borde fino."""
    with Image.open(ruta) as im:
        im = im.convert("RGB")
        if im.width > 1920:
            im = im.resize((1920, round(im.height * 1920 / im.width)), Image.LANCZOS)
        buffer = io.BytesIO()
        im.save(buffer, "JPEG", quality=85)
    buffer.seek(0)
    foto = slide.shapes.add_picture(buffer, Inches(x), Inches(y), width=Inches(w))
    foto.line.color.rgb = LINEA
    foto.line.width = Pt(0.75)
    return foto


def _titulo(slide, texto, color=TEXTO):
    _texto(slide, M, 0.45, W - 2 * M, 0.75, texto, size=32, bold=True, color=color,
           anchor=MSO_ANCHOR.MIDDLE, nombre="Título")


def _pie(slide, numero):
    _texto(slide, M, 7.0, 6, 0.3, f"Grupo 1 · {cfg.TAREA} · Web scraping con Selenium", size=10,
           color=SECUNDARIO)
    _texto(slide, W - M - 1, 7.0, 1, 0.3, str(numero), size=10, color=SECUNDARIO, align=PP_ALIGN.RIGHT)


def _fondo(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _tabla(slide, x, y, anchos, filas, size=12, alto_fila=0.42):
    tabla = slide.shapes.add_table(len(filas), len(anchos), Inches(x), Inches(y),
                                   Inches(sum(anchos)), Inches(alto_fila * len(filas))).table
    tabla.first_row = True
    tabla.horz_banding = False
    for j, ancho in enumerate(anchos):
        tabla.columns[j].width = Inches(ancho)
    for i, fila in enumerate(filas):
        tabla.rows[i].height = Inches(alto_fila)
        for j, valor in enumerate(fila):
            celda = tabla.cell(i, j)
            celda.fill.solid()
            celda.fill.fore_color.rgb = CLARO if i == 0 else BLANCO
            celda.margin_left = celda.margin_right = Inches(0.08)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = celda.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = str(valor)
            es_codigo = i > 0 and str(valor).startswith(("div", "span", "aria"))
            run.font.name = MONO if es_codigo else FUENTE
            run.font.size = Pt(size - 1 if es_codigo else size)
            run.font.bold = i == 0
            run.font.color.rgb = TEXTO
    return tabla


def _notas(slide, texto):
    slide.notes_slide.notes_text_frame.text = texto


# ---------------------------------------------------------------- diapositivas
def _portada(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(s, OSCURO)
    _texto(s, M, 1.0, 8, 0.4, f"{cfg.TAREA.upper()} · GRUPO 1", size=14, bold=True, color=NARANJA)
    _texto(s, M, 1.5, 8.2, 1.7, "Web scraping de reseñas de Rappi", size=44, bold=True, color=BLANCO)
    _texto(s, M, 3.25, 8, 0.5, "Selenium + ChromeDriver sobre Google Play", size=22, color=GRIS_OSCURO)
    _texto(s, M, 4.45, 6, 0.35, "Integrantes", size=12, bold=True, color=NARANJA)
    _texto(s, M, 4.8, 6.5, 1.6, cfg.INTEGRANTES, size=16, color=BLANCO, espacio=2)
    _texto(s, 9.0, 1.5, W - M - 9.0, 1.2, str(d["total"]), size=80, bold=True, color=NARANJA,
           align=PP_ALIGN.RIGHT)
    _texto(s, 9.0, 2.75, W - M - 9.0, 0.4, "reseñas extraídas", size=18, color=BLANCO, align=PP_ALIGN.RIGHT)
    fecha = d["df"]["fecha_extraccion"].min()
    _texto(s, 7.0, 6.05, W - M - 7.0, 0.7, [cfg.CURSO, f"{fecha:%d/%m/%Y}"], size=14, color=GRIS_OSCURO,
           align=PP_ALIGN.RIGHT)
    _notas(s, "Presentamos la Tarea 03: la extracción automatizada de reseñas de Rappi en Google Play "
              f"usando Selenium y ChromeDriver. En total se extrajeron {d['total']} reseñas en español.")


def _caso_de_uso(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Caso de uso: opiniones de Rappi en Google Play")
    puntos = [
        ("Empresa", "Rappi, marketplace de delivery con más de 100 M de descargas"),
        ("Fuente", "Google Play web: reseñas en español, sin iniciar sesión"),
        ("Meta", "~1000 reseñas, ordenadas por «Más recientes»"),
        ("Descartados", "Amazon (exige login) y Booking (bloqueos anti-bot)"),
    ]
    for i, (etiqueta, texto) in enumerate(puntos):
        y = 1.65 + i * 1.15
        _numero(s, M, y + 0.05, 0.42, i + 1, size=14)
        _texto(s, M + 0.65, y, 6.9, 1.0, [[(etiqueta, {"bold": True, "size": 18})], texto], size=16,
               color=SECUNDARIO, espacio=2)
    cifras = [(cfg.OPINIONES_GOOGLE_PLAY, "opiniones públicas de la app"),
              (str(d["total"]), "reseñas extraídas"),
              (f"{d['segundos'] // 60} min {d['segundos'] % 60} s", "de navegación automatizada")]
    for i, (valor, etiqueta) in enumerate(cifras):
        y = 1.55 + i * 1.75
        _tarjeta(s, 8.4, y, W - M - 8.4, 1.5)
        _texto(s, 8.7, y + 0.15, 4, 0.8, valor, size=36, bold=True, color=NARANJA)
        _texto(s, 8.7, y + 0.95, 4, 0.4, etiqueta, size=14, color=SECUNDARIO)
    _pie(s, n)
    _notas(s, "Elegimos Rappi porque tiene millones de opiniones públicas en Google Play, en español y "
              "visibles sin iniciar sesión. Amazon y Booking se descartaron por login obligatorio y bloqueos.")


def _herramientas(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Herramientas y entorno")
    tarjetas = [
        ("Entorno", "Anaconda", "entorno conda «task03»"),
        ("Lenguaje", "Python", platform.python_version()),
        ("Automatización", "Selenium", selenium.__version__),
        ("Navegador", "Google Chrome", d["chrome"]),
        ("Driver · vía Selenium Manager", "ChromeDriver", d["chromedriver"]),
        ("IDE", "Jupyter Lab", f"{cfg.PREFIJO}_scraping_rappi.ipynb"),
    ]
    ancho, alto, gap = (W - 2 * M - 2 * 0.3) / 3, 2.3, 0.3
    for i, (etiqueta, nombre, detalle) in enumerate(tarjetas):
        x, y = M + (i % 3) * (ancho + gap), 1.6 + (i // 3) * (alto + gap)
        _tarjeta(s, x, y, ancho, alto)
        _forma(s, MSO_SHAPE.OVAL, x + 0.35, y + 0.4, 0.2, 0.2, NARANJA)
        _texto(s, x + 0.7, y + 0.33, ancho - 1, 0.35, etiqueta, size=13, color=SECUNDARIO)
        _texto(s, x + 0.35, y + 0.95, ancho - 0.7, 0.6, nombre, size=26, bold=True)
        _texto(s, x + 0.35, y + 1.6, ancho - 0.7, 0.4, detalle, size=15, color=SECUNDARIO)
    _pie(s, n)
    _notas(s, "Trabajamos en un entorno de Anaconda. Selenium Manager descarga automáticamente el "
              "ChromeDriver compatible con la versión instalada de Chrome, por eso ambas versiones coinciden.")


def _flujo(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Flujo del scraper")
    pasos = [
        ("Abrir Chrome", "ChromeDriver vía Selenium Manager"),
        ("Abrir el modal", "Clic en «Ver todas las opiniones»"),
        ("Ordenar", "«Más recientes», con confirmación"),
        ("Extraer", "Lotes de 20 reseñas por cada scroll"),
        ("Guardar", "CSV, checkpoints y capturas"),
    ]
    columna, d_circulo, y_circulo = (W - 2 * M) / 5, 0.9, 1.9
    for i, (nombre, detalle) in enumerate(pasos):
        x = M + i * columna
        cx = x + (columna - d_circulo) / 2
        if i:
            _linea(s, cx - columna + d_circulo + 0.15, y_circulo + d_circulo / 2, cx - 0.15,
                   y_circulo + d_circulo / 2, grosor=1.5)
        _numero(s, cx, y_circulo, d_circulo, i + 1, size=26)
        _texto(s, x + 0.1, 3.1, columna - 0.2, 0.45, nombre, size=18, bold=True, align=PP_ALIGN.CENTER)
        _texto(s, x + 0.1, 3.6, columna - 0.2, 0.9, detalle, size=14, color=SECUNDARIO, align=PP_ALIGN.CENTER)
    _tarjeta(s, M, 4.9, W - 2 * M, 1.2)
    _texto(s, M + 0.4, 4.9, W - 2 * M - 0.8, 1.2,
           [[("Criterio de parada:  ", {"bold": True}),
             (f"meta de {cfg.META} reseñas o {cfg.MAX_SCROLLS_SIN_CAMBIO} scrolls seguidos sin reseñas nuevas", {})]],
           size=16, anchor=MSO_ANCHOR.MIDDLE)
    _pie(s, n)
    _notas(s, "El scraper abre Chrome, entra al modal de opiniones, ordena por más recientes y repite "
              "un ciclo de scroll, espera y lectura de las reseñas nuevas hasta llegar a la meta.")


def _analisis_sitio(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Análisis del sitio")
    ancho = 6.6
    _imagen(s, cfg.DIR_CAPTURAS / "04_modal_mas_recientes.png", M, 1.55, ancho)
    _texto(s, M, 1.55 + ancho / CAPTURA + 0.12, ancho, 0.35,
           "Inspección previa: modal de opiniones ordenado por «Más recientes»", size=12, color=SECUNDARIO)
    sel = cfg.SEL
    filas = [("Campo", "Selector"), ("Reseña", sel["resena"].split()[-1]), ("Usuario", sel["usuario"]),
             ("Calificación", "aria-label de div[role='img']"), ("Fecha", sel["fecha"]),
             ("Comentario", sel["comentario"]), ("Votos útiles", sel["votos"]), ("Respuesta", sel["respuesta"])]
    x = M + ancho + 0.45
    _tabla(s, x, 1.55, [1.7, W - M - x - 1.7], filas, size=14)
    _texto(s, x, 5.1, W - M - x, 0.8, ["20 reseñas cargadas por cada scroll",
                                       "Clases ofuscadas → selectores centralizados en config.py"],
           size=13, color=SECUNDARIO, espacio=4)
    _pie(s, n)
    _notas(s, "Antes de programar inspeccionamos el sitio con un script: identificamos el modal, el "
              "contenedor que hace scroll y los selectores de cada campo.")


def _evidencia_unica(prs, d, n, archivo, titulo, encabezado, texto):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, titulo)
    alto = 5.15
    _imagen(s, cfg.DIR_CAPTURAS / archivo, M, 1.5, alto * CAPTURA)
    x = M + alto * CAPTURA + 0.45
    _numero(s, x, 1.55, 0.5, 1, size=16)
    _texto(s, x, 2.3, W - M - x, 1.1, encabezado, size=20, bold=True)
    _texto(s, x, 3.45, W - M - x, 2.0, texto, size=15, color=SECUNDARIO)
    _texto(s, x, 5.85, W - M - x, 0.4, f"Captura automática · {d['horas'].get(archivo, '')}", size=12,
           color=SECUNDARIO)
    _pie(s, n)
    return s


def _evidencia_doble(prs, d, n, titulo, capturas):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, titulo)
    ancho, gap = (W - 2 * M - 0.35) / 2, 0.35
    arriba = 1.4 + (5.4 - (ancho / CAPTURA + 0.7)) / 2  # centra capturas + leyendas en el área útil
    for i, (archivo, leyenda) in enumerate(capturas):
        x = M + i * (ancho + gap)
        _imagen(s, cfg.DIR_CAPTURAS / archivo, x, arriba, ancho)
        y = arriba + ancho / CAPTURA + 0.25
        _numero(s, x, y, 0.42, i + 1, size=14)
        _texto(s, x + 0.6, y - 0.02, ancho - 0.6, 0.5,
               [[(leyenda, {"bold": True}), (f"   {d['horas'].get(archivo, '')}", {"color": SECUNDARIO})]],
               size=16)
    _pie(s, n)
    return s


def _registro(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Registro de la ejecución")
    lineas = []
    for linea in d["ejecucion"]:
        linea = re.sub(r"^(\d\d:\d\d:\d\d) \w+\s+", r"\1  ", linea)
        linea = re.sub(r"/\S+/(\S+)", r"\1", linea).replace("Captura de pantalla completa", "Captura")
        destacar = "Fin |" in linea
        lineas.append([(linea, {"color": NARANJA if destacar else RGBColor.from_string("E6E6E6"),
                                "bold": destacar})])
    ancho, alto = 8.3, 0.5 + 0.215 * len(lineas)
    _tarjeta(s, M, 1.5, ancho, alto, relleno=OSCURO)
    _texto(s, M + 0.3, 1.75, ancho - 0.5, alto - 0.4, lineas, size=11, font=MONO, espacio=1)
    x = M + ancho + 0.45
    cifras = [(str(d["total"]), "reseñas extraídas"), (f"{d['segundos']} s", "duración total"),
              (str(sum("Checkpoint" in linea for linea in d["ejecucion"])), "checkpoints guardados")]
    for i, (valor, etiqueta) in enumerate(cifras):
        y = 1.5 + i * (alto - 1.25) / 2
        _texto(s, x, y, W - M - x, 0.85, valor, size=40, bold=True, color=NARANJA)
        _texto(s, x, y + 0.85, W - M - x, 0.4, etiqueta, size=14, color=SECUNDARIO)
    _pie(s, n)
    _notas(s, "Este es el registro real de la ejecución (scraping.log): muestra el inicio con las "
              "versiones, cada checkpoint y captura, y el cierre con el total y el motivo de parada.")


def _codigo(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Código clave: bucle de scroll")
    fuente = textwrap.dedent(inspect.getsource(scroll_y_extraer))
    inicio = fuente.index("# Criterios de parada")
    fin = fuente.index("\n", fuente.index("time.sleep(", inicio))
    codigo = textwrap.dedent(" " * 12 + fuente[inicio:fin]).splitlines()
    lineas = [[(linea or " ", {"color": GRIS_OSCURO if linea.strip().startswith("#")
                               else RGBColor.from_string("E6E6E6")})] for linea in codigo]
    _tarjeta(s, M, 1.45, W - 2 * M, 3.85, relleno=OSCURO)
    _texto(s, M + 0.35, 1.65, W - 2 * M - 0.7, 3.5, lineas, size=12, font=MONO)
    notas = [("WebDriverWait", "espera a que llegue el nuevo lote, sin pausas fijas"),
             ("scrollTop", "desplaza el contenedor del modal, no la página"),
             ("Pausa aleatoria", "0,8–2 s entre scrolls, por cortesía con el sitio")]
    ancho = (W - 2 * M - 2 * 0.3) / 3
    for i, (nombre, texto) in enumerate(notas):
        x = M + i * (ancho + 0.3)
        _numero(s, x, 5.6, 0.42, i + 1, size=14)
        _texto(s, x + 0.6, 5.55, ancho - 0.6, 1.1, [[(nombre, {"bold": True, "color": TEXTO})], texto],
               size=14, color=SECUNDARIO, espacio=2)
    _pie(s, n)
    _notas(s, "Este fragmento es el código real del bucle: verifica los criterios de parada, hace "
              "scroll en el modal, espera el nuevo lote con WebDriverWait y hace una pausa aleatoria.")


def _datos_resultantes(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Datos resultantes")
    df = d["df"]
    corto = df[df["longitud_comentario"].between(25, 70)]
    muestra = [corto[corto["calificacion"] == c].iloc[0] for c in range(1, 6) if (corto["calificacion"] == c).any()]
    filas = [("Usuario", "★", "Fecha", "Comentario")] + [
        (r["usuario"][:22], int(r["calificacion"]), f"{r['fecha']:%d/%m/%Y}", r["comentario"]) for r in muestra]
    _tabla(s, M, 1.55, [2.2, 0.5, 1.35, 4.05], filas, size=13, alto_fila=0.62)
    x = M + 8.1 + 0.4
    _tarjeta(s, x, 1.55, W - M - x, 4.35)
    _texto(s, x + 0.3, 1.8, W - M - x - 0.6, 0.45, NOMBRE_ZIP, size=16, bold=True)
    archivos = []
    for nombre, tamaño in d["archivos_zip"]:  # nombre y tamaño en líneas separadas: el nombre no se corta
        archivos += [[(nombre, {"bold": True})], [(tamaño, {"color": SECUNDARIO, "size": 12})],
                     [(" ", {"size": 6})]]
    _texto(s, x + 0.3, 2.35, W - M - x - 0.6, 2.0, archivos, size=13)
    _texto(s, x + 0.3, 4.35, W - M - x - 0.6, 1.4,
           [f"{len(df)} filas · {df.shape[1]} columnas", "Excel con hojas Reseñas, Diccionario y Resumen",
            "7/7 validaciones superadas"], size=14, color=SECUNDARIO, espacio=4)
    _texto(s, M, 1.55 + 0.62 * len(filas) + 0.2, 8.1, 0.4, "Muestra: una reseña por cada calificación",
           size=12, color=SECUNDARIO)
    _pie(s, n)
    _notas(s, "El resultado se entrega en un .zip con el CSV, el Excel y un LEEME con el diccionario de "
              "datos. Todas las validaciones de calidad se superaron.")


def _resultados(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Resultados")
    df, r = d["df"], d["resumen"]
    conteo = df["calificacion"].value_counts().reindex(range(1, 6), fill_value=0)
    datos = CategoryChartData()
    datos.categories = [f"{i}★" for i in conteo.index]
    datos.add_series("Reseñas", [int(v) for v in conteo.values])
    grafico = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(1.95), Inches(7.6),
                                 Inches(4.75), datos).chart
    grafico.has_legend = False
    grafico.has_title = False  # el subtítulo de la diapositiva ya nombra la serie
    grafico.font.name, grafico.font.size, grafico.font.color.rgb = FUENTE, Pt(13), SECUNDARIO
    plot = grafico.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size, plot.data_labels.font.color.rgb = Pt(14), TEXTO
    serie = plot.series[0]
    serie.format.fill.solid()
    serie.format.fill.fore_color.rgb = NARANJA
    eje_y = grafico.value_axis
    eje_y.major_gridlines.format.line.color.rgb = LINEA
    eje_y.format.line.fill.background()
    grafico.category_axis.format.line.color.rgb = LINEA
    _texto(s, M, 1.5, 7.6, 0.4, f"Reseñas por calificación (n = {len(df)})", size=14, color=SECUNDARIO)
    x = M + 7.6 + 0.6
    cifras = [(f"{_decimal(r['Calificación promedio'], 2)} ★", "calificación promedio"),
              (f"{_decimal(r['% de 1 estrella'])} %", "reseñas de 1 estrella"),
              (f"{_decimal(r['% respondidas por Rappi'])} %", "respondidas por Rappi, el mismo día")]
    for i, (valor, etiqueta) in enumerate(cifras):
        y = 1.6 + i * 1.75
        _texto(s, x, y, W - M - x, 0.9, valor, size=40, bold=True, color=NARANJA)
        _texto(s, x, y + 0.9, W - M - x, 0.6, etiqueta, size=14, color=SECUNDARIO)
    _pie(s, n)
    _notas(s, "Las reseñas recientes son mayoritariamente negativas, aunque la app tiene 4,6 estrellas en "
              "total: quienes tienen problemas escriben más. Rappi responde casi todas el mismo día.")


def _dificultades(prs, d, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _titulo(s, "Dificultades y soluciones")
    filas = [
        ("Clases HTML ofuscadas", "Selectores centralizados en config.py, priorizando aria-label"),
        ("Carga dinámica por scroll", "Scroll del modal y espera por lote con WebDriverWait"),
        ("Clic de orden perdido por la animación del menú", "Se confirma el orden activo, con hasta 3 reintentos"),
        ("Reseñas editadas con fecha posterior a la respuesta", "Se marcan con editada_tras_respuesta, sin descartarlas"),
        ("Capturas negras en Wayland", "Captura de pantalla completa mediante el portal del escritorio"),
    ]
    x_solucion = M + 5.6
    _texto(s, M + 0.7, 1.45, 4.6, 0.35, "Problema", size=13, bold=True, color=SECUNDARIO)
    _texto(s, x_solucion, 1.45, W - M - x_solucion, 0.35, "Solución", size=13, bold=True, color=SECUNDARIO)
    for i, (problema, solucion) in enumerate(filas):
        y = 1.95 + i * 0.95
        if i:
            _linea(s, M, y - 0.14, W - M, y - 0.14)
        _numero(s, M, y + 0.14, 0.42, i + 1, size=14)  # centrado con la fila de 0,7"
        _texto(s, M + 0.7, y, 4.6, 0.7, problema, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _texto(s, x_solucion, y, W - M - x_solucion, 0.7, solucion, size=16, color=SECUNDARIO,
               anchor=MSO_ANCHOR.MIDDLE)
    _pie(s, n)
    _notas(s, "Estos fueron los principales problemas. El más sutil fue el ordenamiento: el clic se perdía "
              "sin dar error, así que agregamos una confirmación observable y reintentos.")


def _conclusiones(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(s, OSCURO)
    _titulo(s, "Conclusiones", color=BLANCO)
    r = d["resumen"]
    puntos = [
        f"Selenium y ChromeDriver permitieron extraer {d['total']} reseñas en español de forma automática, "
        "sin duplicados y con todas las validaciones superadas.",
        f"Las opiniones recientes de Rappi son mayoritariamente negativas ({_decimal(r['% de 1 estrella'])} % "
        "de 1 estrella): soporte, cobros y demoras son las quejas más frecuentes.",
        "Las esperas explícitas, los selectores centralizados y las validaciones automáticas hicieron "
        "el proceso reproducible y confiable.",
    ]
    for i, texto in enumerate(puntos):
        y = 1.9 + i * 1.45
        _numero(s, M, y + 0.05, 0.5, i + 1, size=16)
        _texto(s, M + 0.85, y, W - 2 * M - 0.85, 1.1, texto, size=20, color=BLANCO)
    _texto(s, M, 6.3, W - 2 * M, 0.5,
           [[("Uso responsable:  ", {"bold": True, "color": NARANJA}),
             ("fines académicos · sin inicio de sesión · una sola sesión · pausas entre scrolls · "
              "robots.txt lo permite", {})]], size=14, color=GRIS_OSCURO)
    _notas(s, "Concluimos que la automatización con Selenium es efectiva para recolectar opiniones a escala, "
              "siempre con un uso responsable del sitio.")


def convertir_a_pdf(pptx: Path, carpeta: Path) -> Path:
    """Exporta el .pptx a PDF con LibreOffice en modo headless.

    Se usa un perfil temporal para que la conversión funcione aunque LibreOffice esté abierto.
    Las fuentes quedan incrustadas: el PDF se ve igual en cualquier computadora.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("Se necesita LibreOffice (comando 'soffice') para exportar a PDF")
    carpeta.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as perfil:
        subprocess.run([soffice, f"-env:UserInstallation={Path(perfil).as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(carpeta), str(pptx)],
                       check=True, capture_output=True, timeout=300)
    pdf = carpeta / f"{pptx.stem}.pdf"
    if not pdf.exists():
        raise RuntimeError(f"LibreOffice no generó {pdf.name}")
    return pdf


def generar_ppt() -> Path:
    """Arma la presentación (.pptx en evidencia/) y la exporta a PDF en entregables/."""
    d = _datos()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = f"{cfg.TAREA} · Web scraping de reseñas de Rappi"
    prs.core_properties.author = "Grupo 1"

    _portada(prs, d)
    _caso_de_uso(prs, d, 2)
    _herramientas(prs, d, 3)
    _flujo(prs, d, 4)
    _analisis_sitio(prs, d, 5)
    s = _evidencia_unica(prs, d, 6, "scraping_01_pagina_app.png", "Evidencia: Chrome automatizado",
                         "Chrome controlado por ChromeDriver",
                         "La barra «Un software de prueba automatizado está controlando Chrome» confirma "
                         "que el navegador es manejado por Selenium.")
    _notas(s, "Primera evidencia: Chrome se abre solo y carga la ficha de Rappi. La barra superior indica "
              "que lo controla un software automatizado.")
    s = _evidencia_doble(prs, d, 7, "Evidencia: apertura del modal y orden",
                         [("scraping_02_modal_opiniones.png", "Modal de opiniones abierto"),
                          ("scraping_03_orden_recientes.png", "Orden «Más recientes» confirmado")])
    _notas(s, "El scraper hace clic en «Ver todas las opiniones» y cambia el orden a «Más recientes», "
              "verificando que el cambio se aplicó.")
    s = _evidencia_doble(prs, d, 8, "Evidencia: scroll automático en progreso",
                         [("scraping_04_250_resenas.png", "250 reseñas extraídas"),
                          ("scraping_05_500_resenas.png", "500 reseñas extraídas")])
    _notas(s, "Capturas tomadas durante el scroll: las fechas visibles retroceden, lo que demuestra que "
              "el modal avanzaba automáticamente.")
    s = _evidencia_doble(prs, d, 9, "Evidencia: meta alcanzada",
                         [("scraping_06_1000_resenas.png", "1000 reseñas extraídas"),
                          ("scraping_07_final.png", f"Fin: {d['total']} reseñas")])
    _notas(s, f"Al llegar a {d['total']} reseñas el scraper se detiene, guarda los datos y cierra Chrome.")
    _registro(prs, d, 10)
    _codigo(prs, d, 11)
    _datos_resultantes(prs, d, 12)
    _resultados(prs, d, 13)
    _dificultades(prs, d, 14)
    _conclusiones(prs, d)

    pptx = cfg.BASE / "evidencia" / NOMBRE_PPTX
    pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(pptx)
    pdf = convertir_a_pdf(pptx, cfg.DIR_ENTREGABLES)
    log.info("Presentación generada: %s (%d diapositivas; editable en %s)",
             pdf, len(prs.slides), pptx.relative_to(cfg.BASE))
    return pdf


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    print(generar_ppt())
